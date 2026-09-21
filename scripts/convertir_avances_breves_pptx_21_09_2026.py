"""PPTX editable de cinco diapositivas; texto tomado del HTML de avances."""
import json
import hashlib
import re
import zipfile
from pathlib import Path
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'presentaciones/Avances-Resumen-Autocontenido-21-09-2026.html'
DEST = SOURCE.with_suffix('.pptx')
sections = BeautifulSoup(SOURCE.read_text(encoding='utf-8'),'html.parser').select('.slide')
prs = Presentation()
prs.slide_width,prs.slide_height = Pt(960),Pt(540)
prs.core_properties.title = 'Últimos avances de tesis MDS — resumen'
prs.core_properties.author = 'Mario Romero'
prs.core_properties.language = 'es-CL'
INK,MUTED,BLUE,TEAL,SOFT,LINE = '12243B','46556B','003D73','09686B','EFF4F9','CCD7E2'
def txt(node): return node.get_text(' ',strip=True)
def rect(s,x,y,w,h,c):
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    a.fill.solid(); a.fill.fore_color.rgb=RGBColor.from_string(c)
    a.line.fill.background(); a._element.spPr.append(OxmlElement('a:effectLst'))
def text(s,value,x,y,w,h,size=18,bold=False,color=INK):
    a=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h)); f=a.text_frame
    f.word_wrap=True; f.auto_size=MSO_AUTO_SIZE.NONE
    f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.text=line; p.space_before=p.space_after=Pt(0); p.line_spacing=1.12
        p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color)
def base(i):
    section=sections[i]; s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
    rect(s,0,0,960,4,BLUE)
    text(s,txt(section.select_one('.eyebrow')).upper(),50,30,860,20,10,True,MUTED)
    text(s,txt(section.find(['h1','h2'])),50,66,860,79,30 if i==0 else 28,True)
    text(s,txt(section.select_one('.source')),50,489,826,43,8.6,False,MUTED)
    text(s,f'{i+1} / 5',897,515,45,15,9,False,MUTED)
    s.notes_slide.notes_text_frame.text=section.get_text('\n',strip=True)+'\n\nFuente: '+SOURCE.name
    return s,section
def note(s,value,y=427):
    rect(s,50,y,860,49,SOFT);text(s,value,63,y+9,834,35,13.5)

s,section=base(0)
text(s,txt(section.select_one('.lead')),50,162,860,65,23)
for i,a in enumerate(section.select('.metrics-cards article')):
    x=50+i*294; rect(s,x,262,270,2,TEAL)
    text(s,txt(a.b),x,279,270,54,38,True,BLUE)
    text(s,txt(a.p),x,338,255,44,17)
note(s,txt(section.select_one('.note')),420)

s,section=base(1)
text(s,txt(section.figcaption),50,153,860,25,12.5,False,MUTED)
for i,row in enumerate(section.select('.barrow')):
    y=201+i*38
    text(s,txt(row.span),50,y-2,225,29,18,i==0)
    rect(s,284,y,502,20,SOFT)
    width=float(row.select_one('.fill')['style'].split(':')[1].rstrip('%'))
    rect(s,284,y,502*width/100,20,BLUE)
    rect(s,786,y,1,20,MUTED)
    text(s,txt(row.b),807,y-2,102,29,18,True)
text(s,txt(section.select_one('.axis span')),284,391,40,18,10,False,MUTED)
text(s,txt(section.select('.axis span')[1]),703,391,100,18,10,False,MUTED)
note(s,txt(section.select_one('.note')))

s,section=base(2)
text(s,txt(section.caption),50,156,860,28,13,False,MUTED)
xs,widths=[50,286,573],[214,265,337]
for ri,tr in enumerate(section.select('tr')):
    y=209+ri*42
    for ci,cell in enumerate(tr.select('th,td')):
        text(s,txt(cell),xs[ci]+4,y,widths[ci]-8,30,13 if ri==0 else 19,ri==0 or ci==0,MUTED if ri==0 else INK)
    rect(s,50,y+34,860,.6,LINE)
rect(s,50,394,860,73,SOFT);rect(s,50,394,3,73,TEAL)
text(s,txt(section.select_one('.question')),66,409,827,50,20)

s,section=base(3)
for i,a in enumerate(section.select('.findings article')):
    x=50+i*294;rect(s,x,162,272,244,SOFT);rect(s,x,162,3,244,TEAL if i==1 else BLUE)
    text(s,txt(a.h3),x+16,179,238,43,17.5,True)
    text(s,txt(a.select_one('.score')),x+16,235,238,44,32,True,BLUE)
    ps=a.select('p')
    text(s,txt(ps[1]),x+16,293,238,48,17)
    text(s,txt(ps[2]),x+16,347,238,58,13.5,False,MUTED)
note(s,txt(section.select_one('.note')))

s,section=base(4)
for i,li in enumerate(section.select('.nextsteps li')):
    y=167+i*79
    text(s,str(i+1),50,y,32,32,25,True,TEAL)
    text(s,txt(li.strong),95,y,815,29,21,True)
    text(s,txt(li.span),95,y+33,815,32,17.5,False,MUTED)
    rect(s,50,y+66,860,.6,LINE)
note(s,txt(section.select_one('.note')))

prs.save(DEST)
loaded=Presentation(DEST)
report={'slides':len(loaded.slides),'aspect_ratio':'16:9','source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'pptx_sha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),'slides_verified':[]}
def norm(v):return re.sub(r'\s+',' ',v).strip().casefold()
assert len(loaded.slides)==5
for i,(section,s) in enumerate(zip(sections,loaded.slides)):
    visible=norm(' '.join(sh.text for sh in s.shapes if sh.has_text_frame))
    missing=[v for v in section.stripped_strings if norm(v) not in visible]
    assert not missing,(i+1,missing)
    outside=[sh.shape_id for sh in s.shapes if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width or sh.top+sh.height>prs.slide_height]
    assert not outside,(i+1,outside)
    report['slides_verified'].append({'slide':i+1,'source_text_preserved':True,'out_of_bounds':outside})
with zipfile.ZipFile(DEST) as z:
    assert z.testzip() is None
    report['media_files']=[n for n in z.namelist() if n.startswith('ppt/media/')]
    report['external_relationships']=[n for n in z.namelist() if n.endswith('.rels') and b'TargetMode="External"' in z.read(n)]
    assert not report['media_files'] and not report['external_relationships']
(ROOT/'documentacion/Verificacion_PPTX_Avances_Breves_21-09-2026.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(DEST.name)
