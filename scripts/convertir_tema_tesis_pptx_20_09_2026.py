"""Convierte el HTML breve en seis diapositivas 16:9 con objetos editables.

Dependencias: python-pptx, beautifulsoup4. No requiere conexión al ejecutarse.
La fuente del texto es el HTML; no se duplican los objetivos ni los datos.
"""
import hashlib
import json
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'presentaciones/Presentacion-Tema-Tesis-Autocontenida-20-09-2026.html'
DEST = SOURCE.with_suffix('.pptx')
soup = BeautifulSoup(SOURCE.read_text(encoding='utf-8'), 'html.parser')
sections = soup.select('section.slide')
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(960), Pt(540)
prs.core_properties.title = 'Tema de tesis MDS — Mario Romero'
prs.core_properties.author = 'Mario Romero'
prs.core_properties.subject = 'Presentación breve para importar y editar en Canva'
prs.core_properties.language = 'es-CL'
INK, MUTED, BLUE, TEAL, SOFT, LINE = '12243B','46556B','003D73','09686B','EFF4F9','CCD7E2'

def txt(el):
    return el.get_text(' ', strip=True)

def rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape

def text(slide, value, x, y, w, h, size=16, bold=False, color=INK):
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.12
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
    return shape

def panel(slide, title, body, x, y, w, h, accent=BLUE, body_size=16):
    rect(slide, x,y,w,h,SOFT)
    rect(slide, x,y,3,h,accent)
    text(slide,title,x+18,y+16,w-36,29,18,True)
    text(slide,body,x+16,y+54,w-32,h-67,body_size)

def base(section, index):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255,255,255)
    rect(slide,0,0,960,4,BLUE)
    text(slide,txt(section.select_one('.eyebrow')).upper(),50,30,845,20,10,True,MUTED)
    if index:
        text(slide,txt(section.h2),50,64,860,74,28,True)
    text(slide,txt(section.select_one('.source')),50,490,830,40,8.3,False,MUTED)
    text(slide,f'{index+1} / 6',898,510,45,15,9,False,MUTED)
    slide.notes_slide.notes_text_frame.text = (
        f'Fuente: {SOURCE.name}\nRevisión de objetivos y KDD del 22-09-2026.\n\n'
        + section.get_text('\n', strip=True)
    )
    return slide

# 1. Tema y profesor guía.
s = sections[0]; slide = base(s,0)
text(slide,txt(s.h1),50,78,850,165,33,True)
for x,identity in zip([50,280],s.select('.identification div')):
    text(slide,txt(identity.dt),x,271,215,20,11,False,MUTED)
    text(slide,txt(identity.dd),x,294,330,30,18,True)
rect(slide,50,351,860,103,SOFT)
rect(slide,50,351,3,103,TEAL)
text(slide,txt(s.select_one('.label')).upper(),69,367,810,20,10,True,TEAL)
text(slide,txt(s.select_one('.panel p')),69,394,810,49,17)

# 2. Antecedentes.
s = sections[1]; slide = base(s,1)
for i,article in enumerate(s.select('article')):
    panel(slide,txt(article.h3),txt(article.p),50+i*294,161,272,210,TEAL if i==2 else BLUE,16)
rect(slide,50,394,860,71,SOFT)
rect(slide,50,394,3,71,TEAL)
text(slide,txt(s.select_one('.question')),68,407,829,51,18)

# 3. Objetivo general.
s = sections[2]; slide = base(s,2)
rect(slide,50,150,4,176,TEAL)
text(slide,txt(s.select_one('.objective')),72,149,820,182,23)
for x,div in zip([50,513],s.select('.comparison div')):
    label = txt(div.b)
    panel(slide,label,txt(div)[len(label):].strip(),x,365,397,85,BLUE,16)
text(slide,'+',471,390,23,28,20,False,BLUE)

# 4. Los siete objetivos específicos.
s = sections[3]; slide = base(s,3)
for i,li in enumerate(s.select('.goals li')):
    x = 50 if i < 3 else 500
    y = 153 + i*103 if i < 3 else [153,246,327,410][i-3]
    label = txt(li.strong)
    text(slide,str(i+1)+'.',x,y,25,25,17,True,TEAL)
    text(slide,label,x+31,y,368,26,16.5,True)
    text(slide,txt(li)[len(label):].strip(),x+31,y+28,368,70 if i<3 else 61,14.5)

# 5. Metodología.
s = sections[4]; slide = base(s,4)
for i,li in enumerate(s.select('.steps li')):
    x = 50+i*175
    rect(slide,x,148,160,159,SOFT)
    text(slide,f'{i+1:02}',x+12,158,40,22,17,True,TEAL)
    label = txt(li.strong)
    text(slide,label,x+12,186,136,35,13.7,True)
    body_width = 146 if i == 1 else 142
    text(slide,txt(li)[len(label):].strip(),x+(160-body_width)/2,232,body_width,68,11.4)
text(slide,txt(s.caption),50,316,860,22,11.5,False,MUTED)
xs, widths = [50,246,564],[178,301,346]
for ri,tr in enumerate(s.select('tr')):
    y = 345+ri*31
    for ci,cell in enumerate(tr.select('th,td')):
        text(slide,txt(cell),xs[ci]+4,y,widths[ci]-8,29,11 if ri==0 else 13,ri==0 or ci==0,MUTED if ri==0 else INK)
    rect(slide,50,y+30,860,.6,LINE)
rect(slide,50,448,860,36,SOFT)
text(slide,txt(s.select_one('.note')),61,452,838,29,10.7)

# 6. Datos disponibles.
s = sections[5]; slide = base(s,5)
for x,div in zip([50,276,540],s.select('.metrics div')):
    text(slide,txt(div.b),x,148,240,35,28,True,BLUE)
    text(slide,txt(div.span),x,186,320,22,11.5,False,MUTED)
text(slide,txt(s.caption),50,225,860,20,11.5,False,MUTED)
xs,widths = [50,181,626],[125,434,284]
for ri,tr in enumerate(s.select('tr')):
    y = 250+ri*32
    for ci,cell in enumerate(tr.select('th,td')):
        text(slide,txt(cell),xs[ci]+4,y,widths[ci]-8,30,11 if ri==0 else 12,ri==0 or ci==0,MUTED if ri==0 else INK)
    rect(slide,50,y+30,860,.6,LINE)
rect(slide,50,424,860,51,SOFT)
text(slide,txt(s.select_one('.note')),61,434,838,35,12)

prs.save(DEST)

# Verificación del contenido visible y de la estructura editable.
loaded = Presentation(DEST)
def norm(value):
    return re.sub(r'\s+',' ',value).strip()
verification = {'revision_date':'2026-09-22','slides':len(loaded.slides),'aspect_ratio':'16:9','font':'Arial',
                'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
                'pptx_sha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),
                'slides_verified':[]}
assert len(loaded.slides) == 6
for i,(section,slide) in enumerate(zip(sections,loaded.slides)):
    visible = norm(' '.join(sh.text for sh in slide.shapes if sh.has_text_frame))
    missing = [v for v in section.stripped_strings if norm(v).casefold() not in visible.casefold()]
    assert not missing, (i+1,missing)
    outside = [sh.shape_id for sh in slide.shapes if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width or sh.top+sh.height>prs.slide_height]
    assert not outside, (i+1,outside)
    verification['slides_verified'].append({'slide':i+1,'shapes':len(slide.shapes),
        'text_boxes':sum(sh.has_text_frame and bool(sh.text) for sh in slide.shapes),
        'source_text_preserved':True,'out_of_bounds':outside})
with zipfile.ZipFile(DEST) as z:
    assert z.testzip() is None
    verification['media_files'] = [n for n in z.namelist() if n.startswith('ppt/media/')]
    verification['external_relationships'] = [n for n in z.namelist() if n.endswith('.rels') and b'TargetMode="External"' in z.read(n)]
    assert not verification['media_files']
    assert not verification['external_relationships']
(ROOT/'documentacion/Verificacion_PPTX_Tema_22-09-2026.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'Created: {DEST.name}, {DEST.stat().st_size:,} bytes, {len(loaded.slides)} editable slides')
