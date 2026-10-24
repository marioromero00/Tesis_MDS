"""Presentación de avances con el lenguaje visual de la tesis MDS."""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "presentaciones/Avances-Experimentales-Simplificados-24-10-2026.pptx"
AUDIT = ROOT / "documentacion/Verificacion_Avances_Simplificados_24-10-2026.json"
RESULTS = ROOT / "resultados"
W, H, TOTAL = 960, 540, 8
NAVY, BLUE, TEAL = "16324F", "1C5D99", "08777C"
INK, MUTED, PALE, RULE, WHITE = "24394C", "596E80", "EAF1F6", "D7E0E7", "FFFFFF"

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = "Avances experimentales — tesis MDS"
prs.core_properties.subject = "Temporalidad, fusión, robustez y control negativo"
prs.core_properties.author = "Mario Romero"
prs.core_properties.language = "es-CL"


def color(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill):
    item = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    item.fill.solid()
    item.fill.fore_color.rgb = color(fill)
    item.line.fill.background()
    for effect in item._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")
    return item


def dot(slide, x, y, diameter, fill):
    item = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y),
                                  Pt(diameter), Pt(diameter))
    item.fill.solid()
    item.fill.fore_color.rgb = color(fill)
    item.line.fill.background()
    for effect in item._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")
    return item


def trace(slide, x, y, width, phase, amplitude, hue):
    """Trazo esquemático editable; nunca se presenta como registro medido."""
    points = []
    for i in range(43):
        t = i / 42
        drift = 0.3 * math.sin(2 * math.pi * (t + phase))
        pulse = math.sin(2 * math.pi * ((3.5 + phase) * t + phase))
        points.append((x + width * t, y + amplitude * (0.65 * pulse + drift)))
    for (x0, y0), (x1, y1) in zip(points, points[1:]):
        segment = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Pt(x0), Pt(y0), Pt(x1), Pt(y1))
        segment.line.color.rgb = color(hue)
        segment.line.width = Pt(1.2)
        for effect in segment._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")


def text(slide, value, x, y, w, h, size=18, fg=INK, bold=False,
         font="Calibri", align=PP_ALIGN.LEFT):
    item = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = item.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(0)
    frame.margin_top = frame.margin_bottom = Pt(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(value.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        para.alignment = align
        para.space_before = para.space_after = Pt(0)
        para.line_spacing = 1.08
        para.font.name = font
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = color(fg)
    return item


def base(number, eyebrow, title, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(WHITE)
    text(slide, eyebrow, 50, 20, 800, 23, 11, MUTED)
    text(slide, title, 50, 55, 860, 66, 30, NAVY, font="Cambria")
    rect(slide, 50, 484, 860, 1, RULE)
    text(slide, source, 50, 491, 808, 33, 9, MUTED)
    text(slide, str(number), 880, 491, 30, 23, 9, MUTED,
         align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def data_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


# Portada: conserva la jerarquía y los colores de la presentación del tema.
s = base(1, "Universidad de Chile · Magíster en Data Science",
         "Avances experimentales",
         "Tesis MDS · Mario Romero · Toma de Muestra 2 · octubre de 2026.")
text(s, "Temporalidad, fusión y robustez en la predicción de pseudoetiquetas",
     50, 143, 545, 78, 23, NAVY, font="Cambria")
text(s, "Mario Romero", 50, 335, 465, 29, 18, NAVY, True)
text(s, "Profesor guía: Dr. Juan D. Velásquez", 50, 365, 465, 27, 16, INK)
text(s, "Hallazgo principal", 50, 428, 170, 23, 14, TEAL, True)
text(s, "Ridge supera 12 controles temporales en train; en el test ya consultado queda cerca de 1/k.",
     220, 423, 690, 47, 17, INK)
# Esquema de señales. Las trazas son simbólicas, como en la portada previa.
rect(s, 627, 146, 282, 211, PALE)
text(s, "Una sesión, cuatro señales", 642, 155, 249, 26, 16, NAVY, True)
for i, (label, phase, amplitude, hue) in enumerate([
    ("EEG", 0.05, 6, BLUE),
    ("GSR", 0.35, 3.5, TEAL),
    ("Mirada", 0.2, 4.5, BLUE),
    ("Pupila", 0.6, 3, TEAL),
]):
    y = 194 + i * 38
    text(s, label, 642, y, 82, 25, 13, MUTED)
    trace(s, 733, y + 12, 153, phase, amplitude, hue)
notes(s, "Las trazas de la portada son esquemáticas; no son datos medidos. La pseudoetiqueta es una referencia computacional, no una medida directa del estado psicológico.")

# Definición de la tarea, en el mismo lenguaje del esquema profesor/estudiante.
s = base(2, "Diseño del estudio · Toma de Muestra 2",
         "Qué señal crea la etiqueta y cuáles la predicen",
         "Fuente: metodología del manuscrito; Evaluacion_Independiente_y_Validez_24-10-2026.md.")
text(s, "Activación", 50, 156, 160, 31, 20, TEAL, True)
rect(s, 209, 158, 170, 55, PALE)
text(s, "GSR", 221, 165, 145, 41, 21, NAVY, True)
text(s, "→", 397, 166, 36, 35, 25, TEAL, True)
rect(s, 445, 158, 189, 55, WHITE)
text(s, "Pseudoetiqueta", 450, 165, 177, 41, 19, TEAL, True)
text(s, "←", 648, 166, 36, 35, 25, BLUE, True)
rect(s, 697, 158, 210, 55, PALE)
text(s, "EEG · mirada · pupila", 707, 165, 190, 41, 16, NAVY, True)
text(s, "Define el objetivo", 209, 222, 170, 26, 14, MUTED)
text(s, "Intenta predecirlo", 697, 222, 210, 26, 14, MUTED)
rect(s, 50, 288, 860, 1, RULE)
text(s, "La etiqueta de activación clasifica GSR en bajo, medio o alto mediante terciles relativos a cada persona y a toda su sesión.",
     50, 307, 860, 75, 20, INK)
text(s, "Se calcula offline. Atención es un objetivo secundario con una etiqueta distinta.",
     50, 402, 860, 43, 17, MUTED)
notes(s, "La etiqueta de atención depende de mirada y pupila; para cada objetivo, las señales que crean la etiqueta quedan fuera del predictor correspondiente. Los terciles no son un umbral fisiológico común ni una etiqueta causal en tiempo real.")

# Muestra: la partición reproduce el lenguaje de la diapositiva original.
s = base(3, "Toma de Muestra 2 · navegación web en laboratorio",
         "La evaluación separa a las personas",
         "Fuentes: partición de participantes; Evaluacion_Independiente_y_Validez_24-10-2026.md.")
text(s, "48 participantes registrados", 50, 153, 490, 40, 23, NAVY, True)
text(s, "↓", 52, 202, 42, 35, 23, TEAL)
text(s, "41 en el conjunto principal", 94, 199, 445, 41, 21, TEAL, True)
text(s, "6 excluidos por calidad o falta de señal · 1 reservado para sensibilidad",
     94, 247, 790, 34, 16, MUTED)
text(s, "Partición por persona", 50, 299, 300, 28, 17, NAVY, True)
for x, w, count, label, hue in [
    (50, 512, "25", "train · 13.524 ventanas", BLUE),
    (566, 166, "8", "validation", TEAL),
    (736, 174, "8", "test · 4.025 ventanas", NAVY),
]:
    rect(s, x, 337, w, 65, hue)
    text(s, count, x + 12, 342, 55, 40, 25, WHITE, True)
    text(s, label, x + 12, 378, w - 22, 19, 12, WHITE)
text(s, "BA macro: recall medio de las clases presentes por persona, promediado entre personas.",
     50, 420, 860, 49, 17, INK)
notes(s, "Train usa cinco folds: cada persona se predice fuera del fold de ajuste. Validation ya se usó para comparar entradas. Test ya se había consultado, así que su evaluación adicional no confirma generalización independiente. Los intervalos descriptivos remuestrean personas, no ventanas.")

# Resultados factoriales. Tabla y intervalos descriptivos editables.
s = base(4, "Contraste temporalidad × fusión · train",
         "Las diferencias medias son pequeñas",
         "Fuente: Contraste_Temporalidad_Fusion_24-10-2026.md; 25 personas, 5 folds, 2 semillas.")
text(s, "BA macro por persona", 50, 148, 378, 27, 16, NAVY, True)
text(s, "Ventana repetida", 161, 184, 165, 25, 13, MUTED)
text(s, "Historial real", 329, 184, 130, 25, 13, MUTED)
rect(s, 50, 215, 420, 1, RULE)
for y, name, c, r in [
    (225, "Temprana", "0,3524", "0,3612"),
    (285, "Tardía", "0,3596", "0,3630"),
]:
    text(s, name, 50, y, 104, 40, 16, INK)
    text(s, c, 161, y, 145, 40, 20, NAVY, True)
    text(s, r, 329, y, 130, 40, 20, NAVY, True)
    rect(s, 50, y + 51, 420, 1, RULE)
text(s, "Contrastes en puntos porcentuales", 517, 148, 389, 27, 16, NAVY, True)
rows = {r["contrast"]: r for r in data_rows(
    RESULTS / "contraste_temporalidad_fusion_24-10-2026/contrastes.csv")
    if r["task"] == "arousal"}
plot_x, plot_w, plot_lo, plot_hi = 550, 326, -1.5, 2.5
map_x = lambda value: plot_x + (value - plot_lo) / (plot_hi - plot_lo) * plot_w
for tick in [-1, 0, 1, 2]:
    x = map_x(tick)
    rect(s, x, 215, 1, 157, NAVY if tick == 0 else RULE)
    text(s, f"{tick:+d}" if tick else "0", x - 16, 378, 34, 22,
         11, MUTED, align=PP_ALIGN.CENTER)
for y, key, label in [
    (246, "temporal_early", "Historial E"),
    (321, "fusion_history", "Fusión con R"),
]:
    row = rows[key]
    lo, delta, hi = [float(row[k]) * 100 for k in ("ci_low", "delta_ba", "ci_high")]
    text(s, label, 517, y - 23, 210, 21, 12, MUTED)
    rect(s, map_x(lo), y, map_x(hi) - map_x(lo), 2, TEAL)
    dot(s, map_x(delta) - 5, y - 4, 10, TEAL)
    text(s, f"{delta:+.2f} pp".replace(".", ","), 792, y - 27,
         100, 24, 14, NAVY, True, align=PP_ALIGN.RIGHT)
text(s, "IC descriptivos por remuestreo de personas; ambos cruzan cero.",
     50, 428, 860, 39, 16, INK)
notes(s, "E: fusión temprana; T: fusión tardía. C: repetir la ventana actual; R: historial real. La fusión tardía usa expertos por modalidad y otra capacidad. El historial alcanza hasta nueve segundos. No se corrige la búsqueda previa ni comparaciones múltiples.")

# Diagnóstico por clase: gráfico de barras a partir del CSV, con eje visible.
s = base(5, "Diagnóstico del contraste · activación",
         "Con fusión tardía cae el recall de la clase media",
         "Fuente: resultados/contraste_temporalidad_fusion_24-10-2026/diagnostico_clases.csv.")
class_rows = {(r["scope"], r["class"]): float(r["recall"]) * 100
              for r in data_rows(RESULTS / "contraste_temporalidad_fusion_24-10-2026/diagnostico_clases.csv")
              if r["task"] == "arousal"}
text(s, "Recall por clase · escala 0–60 %", 50, 145, 500, 28, 16, NAVY, True)
text(s, "E-R", 620, 146, 70, 25, 13, BLUE, True)
text(s, "T-R", 713, 146, 70, 25, 13, TEAL, True)
text(s, "E-R: temprana con historial · T-R: tardía con historial",
     50, 173, 650, 22, 13, MUTED)
for i, label in enumerate(["bajo", "medio", "alto"]):
    y = 203 + i * 79
    a = class_rows[("E-R", label)]
    b = class_rows[("T-R", label)]
    text(s, label.capitalize(), 50, y + 8, 101, 35, 18, INK, True)
    rect(s, 170, y, 410, 13, PALE)
    rect(s, 170, y, 410 * a / 60, 13, BLUE)
    rect(s, 170, y + 22, 410, 13, PALE)
    rect(s, 170, y + 22, 410 * b / 60, 13, TEAL)
    text(s, f"{a:.1f} %".replace(".", ","), 620, y - 3, 87, 23, 16, NAVY)
    text(s, f"{b:.1f} %".replace(".", ","), 713, y + 18, 87, 23, 16, NAVY)
text(s, "Con T-R la clase media cae de 23,4 % a 8,2 %, aunque la BA global cambia poco.",
     50, 448, 860, 30, 17, INK)
notes(s, "Las barras muestran E-R y T-R, en ese orden, para activación. BA macro: E-R 0,3612 y T-R 0,3630. Para atención, objetivo secundario, historial temprano baja BA 0,65 pp y fusión tardía con historial sube 1,01 pp; resultados exploratorios. No verifican estados psicológicos.")

# Test: se evita un gráfico que exageraría una diferencia de 0,0015.
s = base(6, "Ridge de activación · evaluación adicional",
         "En test, Ridge queda cerca de 1/k",
         "Fuentes: Evaluacion_Independiente_y_Validez_24-10-2026.md; Robustez_Sin_Cohorte_Nueva_24-10-2026.md.")
text(s, "BA macro por persona", 50, 157, 860, 27, 16, NAVY, True)
for y, label, number in [
    (205, "Train · fuera de fold", "0,3843"),
    (271, "Test · ensamble fijo", "0,3557"),
    (337, "Referencia 1/k · test", "0,3542"),
]:
    rect(s, 50, y - 5, 860, 1, RULE)
    text(s, label, 50, y + 1, 600, 48, 20, INK)
    text(s, number, 680, y + 1, 229, 48, 28, NAVY, True,
         align=PP_ALIGN.RIGHT)
rect(s, 50, 404, 860, 1, RULE)
text(s, "4 de 8 personas superan 1/k. Recall de la clase media en test: 21,8 %.",
     50, 418, 860, 54, 18, TEAL, True)
notes(s, "1/k es la BA de predecir siempre una clase presente; en test promedia 0,3542 porque una persona tiene dos clases. Train usa un modelo por fold; test promedia probabilidades de cinco modelos. El test ya se había consultado; la diferencia es descriptiva y no confirma desempeño independiente.")

# Control negativo: los doce puntos se leen directamente desde el CSV.
s = base(7, "Control negativo temporal · train",
         "Desplazar la etiqueta GSR reduce la BA",
         "Fuente: resultados/control_negativo_temporal_24-10-2026/metricas_repeticiones.csv.")
controls = data_rows(RESULTS / "control_negativo_temporal_24-10-2026/metricas_repeticiones.csv")
assert len(controls) == 12
xs, ys, pw, ph = 112, 169, 570, 223
text(s, "BA macro por persona · eje 0,30–0,40", 50, 141, 570, 23,
     13, MUTED)
low, high = .30, .40
map_y = lambda value: ys + (high - value) / (high - low) * ph
for tick in [.30, .32, .34, .36, .38, .40]:
    y = map_y(tick)
    rect(s, xs, y, pw, 1, RULE)
    text(s, f"{tick:.2f}".replace(".", ","), 50, y - 11, 53, 23,
         12, MUTED, align=PP_ALIGN.RIGHT)
original = .3843
rect(s, xs, map_y(original), pw, 2, NAVY)
for i, row in enumerate(controls):
    x = xs + 17 + i * (pw - 34) / 11
    dot(s, x - 5, map_y(float(row["ba_macro"])) - 5, 10, TEAL)
    text(s, str(i + 1), x - 15, 405, 30, 22, 11, MUTED,
         align=PP_ALIGN.CENTER)
text(s, "Repetición del control", 236, 432, 300, 25, 13, MUTED,
     align=PP_ALIGN.CENTER)
text(s, "Original", 712, 182, 194, 25, 14, NAVY, True)
text(s, "0,3843", 712, 211, 194, 44, 28, NAVY, True)
text(s, "Control, media", 712, 278, 194, 25, 14, TEAL, True)
text(s, "0,3268", 712, 307, 194, 44, 28, TEAL, True)
text(s, "Diferencia +5,76 pp", 712, 389, 194, 46, 16, INK, True)
notes(s, "Se desplazó circularmente la etiqueta GSR dentro de cada registro al ajustar el modelo, preservando las frecuencias de clase. Ninguno de los 12 controles alcanzó el original. El modelo original se había seleccionado con etiquetas de train. En test, GSR 2 s cambia 19,7 % de etiquetas y PCA GSR 30,5 %; las diferencias de BA incluyen cero. No es prueba de validez psicológica ni de generalización a personas nuevas.")

# Balance: cierra con una afirmación, su límite y un siguiente experimento.
s = base(8, "Balance del trabajo experimental",
         "Qué muestran los experimentos y qué falta",
         "Fuentes: informes experimentales del 24-10-2026; Evaluacion_Independiente_y_Validez_24-10-2026.md.")
text(s, "Observado", 50, 161, 189, 28, 16, TEAL, True)
text(s, "En train, Ridge depende de la alineación temporal con la pseudoetiqueta GSR.",
     252, 153, 658, 73, 22, NAVY, font="Cambria")
rect(s, 50, 239, 860, 1, RULE)
text(s, "Límite", 50, 262, 189, 28, 16, TEAL, True)
text(s, "No se confirma una ventaja general del historial o la fusión. Test ya se había consultado y falta una medida externa del estado psicológico.",
     252, 254, 658, 93, 19, INK)
rect(s, 50, 365, 860, 1, RULE)
text(s, "Para confirmar", 50, 388, 189, 28, 16, TEAL, True)
text(s, "Fijar modelo y protocolo antes de evaluar personas nuevas; contrastar la etiqueta con una medida independiente.",
     252, 382, 658, 77, 19, INK)
notes(s, "Conservar la distinción entre predicción de una pseudoetiqueta y validez del constructo. La disponibilidad de una cohorte externa sigue pendiente de confirmación.")

prs.save(DEST)
opened = Presentation(DEST)
assert len(opened.slides) == TOTAL
report = {"slides": TOTAL, "size_points": [W, H], "visible_text": [], "out_of_bounds": []}
for index, slide in enumerate(opened.slides, 1):
    values = [item.text for item in slide.shapes if item.has_text_frame and item.text.strip()]
    report["visible_text"].append({"slide": index, "items": len(values),
                                   "characters": sum(map(len, values))})
    for item in slide.shapes:
        if (item.left < 0 or item.top < 0 or
                item.left + item.width > opened.slide_width or
                item.top + item.height > opened.slide_height):
            report["out_of_bounds"].append({"slide": index, "shape_id": item.shape_id})
assert not report["out_of_bounds"]
AUDIT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Presentación creada: {DEST}")
