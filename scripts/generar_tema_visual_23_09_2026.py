"""HTML sin dependencias y PPTX editable a partir de una misma composición.

Los gráficos son esquemas explicativos, no señales ni resultados experimentales.
La redacción formal de objetivos se lee de las notas de tesis.
"""
import hashlib
import html
import json
import math
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'presentaciones/Presentacion-Tema-Tesis-Autocontenida-20-09-2026'
NOTES = ROOT.parent / '01 Tesis MDS/Escritura/Estructura Tesis.md'
md = NOTES.read_text(encoding='utf-8')
objective = md.split('#### Objetivo General\n', 1)[1].split('\n\n', 1)[0]
goals = md.split('#### Objetivos Específicos\n', 1)[1].split('\n\n', 1)[0]
goals = [re.sub(r'^\d+\. ', '', s) for s in goals.splitlines()]
assert len(goals) == 7

INK, MUTED, BLUE, TEAL, LIGHT, RULE, WHITE = (
    '16324F', '4B5968', '145DA0', '066A70', 'F1F5F8', 'CDD7DF', 'FFFFFF')
slides = []

def slide(title, topic, note):
    s = {'title': title, 'topic': topic, 'note': note, 'objects': []}
    slides.append(s)
    rect(s, 46, 41, 34, 4, TEAL)
    text(s, topic, 92, 27, 1080, 27, 15, color=MUTED)
    text(s, title, 46, 70, 1180, 76, 38, font='Cambria', tag='h2')
    line(s, 46, 667, 1234, 667, RULE, 1)
    # Footnotes are appended last, keeping the reading order of the body intact.
    return s

def text(s, value, x, y, w, h, size=22, bold=False, color=INK, font='Arial', tag='p', cls='', mobile_after=''):
    s['objects'].append(dict(kind='text', value=value, x=x, y=y, w=w, h=h,
                             size=size, bold=bold, color=color, font=font, tag=tag, cls=cls, mobile_after=mobile_after))

def rect(s, x, y, w, h, fill=LIGHT, stroke=None, radius=False):
    s['objects'].append(dict(kind='rect', x=x, y=y, w=w, h=h,
                             fill=fill, stroke=stroke, radius=radius))

def line(s, x1, y1, x2, y2, color=TEAL, width=2):
    s['objects'].append(dict(kind='line', x1=x1, y1=y1, x2=x2, y2=y2, color=color, width=width))

def arrow(s, x1, y1, x2, y2, color=TEAL):
    line(s, x1, y1, x2, y2, color, 2)
    angle = math.atan2(y2-y1, x2-x1)
    for side in (-1, 1):
        a = angle + side*0.55
        line(s, x2, y2, x2-10*math.cos(a), y2-10*math.sin(a), color, 2)

def box(s, value, x, y, w, h, size=20, fill=LIGHT, color=INK):
    rect(s, x, y, w, h, fill)
    text(s, value, x+14, y+12, w-28, h-18, size, color=color, cls='diagram-token')

def finish(s):
    text(s, s['note'], 46, 678, 1135, 30, 11.5, color=MUTED, cls='source')
    text(s, f'{len(slides)}', 1200, 678, 32, 24, 13, color=MUTED, cls='page-number')

# 1. Tema y antecedentes. El dibujo muestra el problema temporal sin simular resultados.
s = slide('Modelos temporales multimodales para la detección de atención y activación emocional durante la navegación web',
          'Universidad de Chile · Magíster en Data Science',
          'Fondecyt 1231122 · Toma de Muestra 2 · 23 de septiembre de 2026. Esquema ilustrativo; no representa datos observados.')
heading = s['objects'][2]
heading.update(w=738, h=219, size=37, tag='h1')
text(s, 'Mario Romero', 46, 315, 680, 38, 26, True)
text(s, 'Profesor guía: Dr. Juan D. Velásquez', 46, 360, 715, 35, 23)
text(s, 'Antecedentes', 46, 437, 680, 32, 22, True, tag='h2')
text(s, 'El proyecto registra señales fisiológicas mientras las personas navegan por la web. La tesis estudia si su evolución en el tiempo ayuda a predecir atención y activación.',
     46, 482, 680, 115, 23)
rect(s, 831, 92, 395, 342, LIGHT)
text(s, 'Una sesión, cuatro señales', 850, 112, 350, 35, 22, True)
for i, label in enumerate(['EEG', 'GSR', 'Mirada', 'Pupila']):
    y = 188+i*54
    text(s, label, 850, y-11, 85, 25, 16, color=MUTED)
    pts=[]
    for j in range(41):
        v = math.sin(j*(.76 if i==0 else .24)+i)*9 + math.sin(j*.11)*5
        pts.append((947+j*6, y+v))
    for a,b in zip(pts,pts[1:]): line(s,*a,*b, BLUE if i%2==0 else TEAL,1.7)
arrow(s, 950, 412, 1197, 412, MUTED)
text(s, 'Tiempo', 1050, 439, 144, 27, 16, color=MUTED)
text(s, 'Pregunta de la tesis', 831, 491, 390, 30, 20, True, color=TEAL)
text(s, '¿Aporta información el historial de las señales?', 831, 532, 390, 90, 27, font='Cambria')
finish(s)

# 2. Datos. Conteos auditados y significado de las cuatro modalidades.
s=slide('Datos disponibles', 'Toma de Muestra 2 · Navegación web en laboratorio',
        'Fuente: preprocesamiento y partición guardada. 25.992 ventanas generadas; 22.201 válidas antes de separar el caso de sensibilidad.')
text(s,'48 participantes registrados',46,158,690,42,29,True)
arrow(s,64,235,64,281)
text(s,'41 en el conjunto principal',95,241,620,38,26,True,color=TEAL)
text(s,'6 excluidos por calidad o falta de señal\n1 analizado por separado por sincronización',95,291,650,66,21,color=MUTED)
text(s,'La partición se hace por persona',46,384,690,32,21,True)
for x,w,n,label in [(46,381,'25','Entrenamiento'),(434,122,'8','Validación'),(563,122,'8','Prueba')]:
    rect(s,x,434,w,88,BLUE if n=='25' else TEAL)
    text(s,n,x+15,445,w-26,38,28,True,color=WHITE)
    text(s,label,x+15,487,w-26,28,14,color=WHITE)
text(s,'Cada ventana resume 2 segundos de registro.\nLa siguiente comienza 1 segundo después.',46,559,661,72,23)
for i,(name,desc) in enumerate([
    ('EEG · actividad eléctrica cerebral','16 canales · 125 Hz'),
    ('GSR · respuesta electrodérmica','Cambios en la conductancia de la piel · ≈15 Hz'),
    ('Eye tracking · seguimiento ocular','Posición de la mirada y fijaciones · ≈60 Hz'),
    ('Pupilometría','Diámetro de la pupila y validez de la medición')]):
    y=165+i*115
    line(s,768,y+2,768,y+76,TEAL,3)
    text(s,name,791,y,431,47,22,True,tag='h3')
    text(s,desc,791,y+53,431,52,19,color=MUTED)
finish(s)

# 3. Objetivo general y contraste temporal.
s=slide('Objetivo general', 'Temporalidad y fusión de señales',
        'Las comparaciones usan las mismas etiquetas y particiones. Esquema de entradas; no representa predicciones ni resultados.')
text(s,objective,46,157,1183,135,26,cls='objective')
text(s,'Modelo estático',46,343,530,36,25,True,tag='h3')
box(s,'Ventana actual',46,405,224,62)
arrow(s,283,436,346,436)
box(s,'Predicción',359,405,217,62)
text(s,'Usa las características de una ventana.',46,493,540,56,22,color=MUTED)
text(s,'Modelo temporal',705,343,530,36,25,True,tag='h3')
for i,label in enumerate(['t−2','t−1','t']): box(s,label,705+i*76,405,67,62,20)
arrow(s,937,436,986,436)
box(s,'Predicción',998,405,224,62)
text(s,'Usa varias ventanas en su orden temporal.',705,493,515,60,22,color=MUTED)
text(s,'En ambos casos se comparan distintas formas de combinar las señales: fusión temprana, intermedia y tardía.',
     46,582,1174,60,23)
finish(s)

# 4. Objetivos específicos íntegros: presentación editorial, sin siete tarjetas.
s=slide('Objetivos específicos', 'Qué se hará en la tesis',
        'Ablación: retirar una señal o el historial y medir cuánto cambia el desempeño. El análisis de fatiga y habituación es descriptivo.')
for i,g in enumerate(goals):
    if i<3: x,y,w=46,165+i*151,557
    else: x,y,w=680,155+(i-3)*127,551
    text(s,str(i+1),x,y,38,35,24,True,color=TEAL)
    text(s,g,x+51,y,w-51,112,21,cls='goal')
finish(s)

# 5. KDD como recorrido. Cada fila explica acción y producto de esa etapa.
s=slide('Metodología KDD', 'Knowledge Discovery in Databases · Descubrimiento de conocimiento en datos',
        'Adaptación de Fayyad, Piatetsky-Shapiro y Smyth (1996), AI Magazine 17(3), 37–54. DOI: 10.1609/aimag.v17i3.1230.')
text(s,'KDD organiza el trabajo desde los registros originales hasta la interpretación de los resultados.',46,150,1186,55,24)
text(s,'Etapa',92,219,285,30,18,True,color=MUTED,cls='diagram-token')
text(s,'Qué se hace en esta tesis',418,219,406,30,18,True,color=MUTED,cls='diagram-token')
text(s,'Qué se obtiene',931,219,295,30,18,True,color=MUTED,cls='diagram-token')
stages=[
    ('Selección','Elegir participantes, sesiones y señales de la Toma de Muestra 2.','Corpus de trabajo'),
    ('Preprocesamiento','Revisar calidad, tratar faltantes y alinear las señales en el tiempo.','Registros sincronizados'),
    ('Transformación','Crear ventanas de 2 s, características, pseudoetiquetas y secuencias.','Datos para los modelos'),
    ('Minería de datos','Entrenar modelos estáticos y temporales; ajustar sus parámetros.','Modelos ajustados'),
    ('Interpretación y evaluación','Comparar el desempeño y revisar qué señales e historial aportan.','Resultados y limitaciones')]
for i,(stage,action,output) in enumerate(stages):
    y=263+i*66
    if i%2==0:rect(s,46,y-3,1188,63,LIGHT)
    text(s,str(i+1),59,y+8,31,31,20,True,color=TEAL,cls='diagram-token')
    text(s,stage,100,y+5,293,52,21,True,tag='h3',cls='kdd-stage')
    text(s,action,418,y+5,421,57,19,cls='kdd-action')
    arrow(s,863,y+27,908,y+27)
    text(s,output,931,y+9,291,48,20,color=TEAL,cls='kdd-output')
    if i<4:arrow(s,74,y+44,74,y+62,MUTED)
arrow(s,817,631,111,631,TEAL)
text(s,'Se vuelve a las etapas anteriores cuando hace falta, usando solo los datos de desarrollo.',
     111,603,1098,29,18,color=MUTED)
finish(s)

# 6. Dos rutas de etiquetado y predicción, sin circularidad.
s=slide('Cómo se construyen las etiquetas', 'Transformación · Etapa 3 de KDD',
        'Eye tracking mide la mirada, no la atención directa. GSR aproxima activación (arousal), sin identificar alegría, tristeza u otras emociones.')
text(s,'Una pseudoetiqueta es un valor estimado a partir de las señales: se usa como referencia para entrenar y evaluar.',46,156,1182,64,25)
text(s,'Modelo Profesor: crea la referencia',46,250,639,32,21,True,color=MUTED,cls='diagram-token')
text(s,'Modelo Estudiante: intenta predecirla',727,250,506,32,21,True,color=MUTED,cls='diagram-token')
for i,(target,teacher,student) in enumerate([
    ('Atención','Mirada + pupilometría','EEG + GSR'),
    ('Activación','GSR','EEG + mirada + pupilometría')]):
    y=309+i*145
    text(s,target,46,y,164,36,25,True,color=TEAL,tag='h3',mobile_after=f'El Modelo Profesor crea la pseudoetiqueta con {teacher}. El Modelo Estudiante intenta predecirla usando {student}.')
    box(s,teacher,225,y-1,291,77,21)
    arrow(s,526,y+37,573,y+37)
    rect(s,584,y-1,182,77,WHITE,TEAL)
    text(s,'Pseudoetiqueta',597,y+21,158,42,19,True,color=TEAL,cls='diagram-token')
    arrow(s,1002,y+37,781,y+37,BLUE)
    box(s,student,1011,y-1,223,77,20)
    text(s,'predice',846,y+4,138,26,16,color=BLUE,cls='diagram-token')
text(s,'Las señales que crean la etiqueta quedan fuera de las entradas del modelo que la predice.',46,599,1186,49,24,True)
finish(s)

# 7. Fusión y evaluación. Diagramas muestran el lugar de combinación.
s=slide('Cómo se comparan los modelos', 'Minería de datos y evaluación · Etapas 4 y 5 de KDD',
        'Evaluación offline; búsquedas previas exploratorias. Recall: proporción detectada por clase. BA: promedio de esos recalls. Macro-F1: promedio por clase del equilibrio entre precisión y recall.')
text(s,'La fusión cambia el momento en que se combinan las señales.',46,155,1187,47,25)
for x,title,desc in [(46,'Temprana','Unir características antes del modelo.'),
                     (458,'Intermedia','Unir representaciones aprendidas.'),
                     (870,'Tardía','Combinar las predicciones de cada modelo.')]:
    text(s,title,x,219,362,36,25,True,color=TEAL,tag='h3')
    text(s,desc,x,265,362,58,20)
    if title=='Temprana':
        box(s,'A',x,353,58,43,17);box(s,'B',x,416,58,43,17)
        arrow(s,x+65,375,x+113,404);arrow(s,x+65,438,x+113,404)
        box(s,'A + B',x+123,384,91,45,17)
        arrow(s,x+222,406,x+248,406);box(s,'Modelo',x+259,380,113,54,17)
    elif title=='Intermedia':
        box(s,'A',x,353,52,43,17);box(s,'B',x,416,52,43,17)
        for y in (375,438):
            arrow(s,x+60,y,x+83,y);box(s,'Red',x+93,y-22,72,45,17)
        arrow(s,x+173,375,x+199,403);arrow(s,x+173,438,x+199,403)
        box(s,'Unión',x+209,380,155,54,17)
    else:
        for y,label in [(375,'Modelo A'),(438,'Modelo B')]:
            box(s,label,x,y-22,155,45,17);arrow(s,x+164,y,x+197,406)
        box(s,'Promedio /\ncombinación',x+211,371,156,75,17)
text(s,'A y B representan modalidades, como EEG y GSR. Solo se usan las señales permitidas para cada variable.',46,481,1183,44,18,color=MUTED)
line(s,46,543,1234,543,RULE,1)
text(s,'Comparación justa',46,568,303,38,24,True)
text(s,'Mismas etiquetas y participantes.\nAjuste dentro del entrenamiento; evaluación final reservada.',379,563,437,80,19)
text(s,'Exactitud balanceada (BA),\nmacro-F1 y recall por clase.\nVariación entre participantes.',883,560,341,90,19)
finish(s)

# HTML: diagramas vectoriales incluidos. A 760 px el texto adopta orden de lectura.
css='''
:root{--ink:#16324f;--muted:#4b5968;--teal:#066a70;--bg:#e8eef3}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:18px/1.4 Arial,sans-serif}
.toolbar{max-width:1280px;margin:auto;padding:12px 0;display:flex;gap:9px;align-items:center;flex-wrap:wrap}
button,select{font:inherit;color:var(--ink);background:white;border:1px solid #4b5968;border-radius:3px;min-height:44px;padding:7px 13px}button{cursor:pointer}button:disabled{opacity:.45;cursor:default}select{max-width:310px}
:focus-visible{outline:3px solid var(--teal);outline-offset:4px}.skip{position:absolute;top:-80px;left:12px;z-index:10;background:white;padding:10px}.skip:focus{top:10px;color:var(--ink)}
main{max-width:1280px;margin:auto}.slide{background:white;position:relative;aspect-ratio:16/9;container-type:inline-size;margin-bottom:24px;isolation:isolate}.slide[hidden],.toolbar[hidden]{display:none}
.t{position:absolute;margin:0;line-height:1.17;font-weight:400;white-space:pre-line;overflow-wrap:normal}.graphic{position:absolute;inset:0;width:100%;height:100%;z-index:-1;pointer-events:none}
.status{margin-left:auto;font-size:15px}.footer{max-width:1280px;margin:auto;padding:0 0 18px;font-size:14px;color:var(--muted)}.mobile-only{display:none}
@media(min-width:761px) and (max-width:1312px){main,.toolbar,.footer{margin-left:16px;margin-right:16px}}
@media(max-width:760px){.toolbar{padding:10px;gap:8px;font-size:15px}select{max-width:250px}.status{margin-left:0}main{padding:0 10px}.slide{aspect-ratio:auto;container-type:normal;padding:24px;margin-bottom:18px}.graphic{display:none}.t{position:static!important;width:auto!important;height:auto!important;font-size:18px!important;line-height:1.45!important;margin:0 0 16px!important;color:var(--ink)!important}h1.t{font-size:29px!important}h2.t{font-size:27px!important}h3.t{font-size:22px!important;margin-top:24px!important}.source{font-size:13px!important;border-top:1px solid #cdd7df;padding-top:16px}.page-number,.diagram-token{display:none}.mobile-only{display:block;font-size:18px;line-height:1.45}.kdd-output:before{content:'Resultado: ';font-weight:bold}.footer{padding:0 15px 18px}}
@media(prefers-reduced-motion:reduce){*{scroll-behavior:auto!important}}
@media print{@page{size:A4 landscape;margin:7mm}body{background:white}.toolbar,.footer,.skip{display:none!important}main{max-width:none}.slide,.slide[hidden]{display:block!important;margin:0;break-after:page;break-inside:avoid;aspect-ratio:16/9}.slide:last-child{break-after:auto}}
'''
nav='''<a class="skip" href="#contenido">Saltar al contenido</a><nav class="toolbar" aria-label="Controles de presentación" hidden>
<button id="prev" type="button" aria-label="Diapositiva anterior">← Anterior</button><button id="next" type="button" aria-label="Diapositiva siguiente">Siguiente →</button>
<label for="jump">Ir a</label><select id="jump"></select><button id="all" type="button" aria-pressed="false">Ver todas</button><button id="print" type="button">Imprimir / PDF</button><span id="status" class="status" role="status" aria-live="polite"></span></nav>'''
js='''(()=>{'use strict';const slides=[...document.querySelectorAll('.slide')],jump=document.querySelector('#jump'),status=document.querySelector('#status'),prev=document.querySelector('#prev'),next=document.querySelector('#next'),all=document.querySelector('#all');let index=0,list=false;
function fromHash(){const n=Number(location.hash.slice(1));return Number.isInteger(n)&&n>=1&&n<=slides.length?n-1:0}index=fromHash();
slides.forEach((s,i)=>{const o=document.createElement('option');o.value=i;o.textContent=`${i+1}. ${s.dataset.title}`;jump.append(o)});
function render(){slides.forEach((s,i)=>s.hidden=!list&&i!==index);jump.value=index;prev.disabled=list||index===0;next.disabled=list||index===slides.length-1;status.textContent=list?`${slides.length} diapositivas`:`${index+1} / ${slides.length}`;all.setAttribute('aria-pressed',String(list));all.textContent=list?'Ver una':'Ver todas';history.replaceState(null,'',`#${index+1}`)}
function go(i){index=Math.max(0,Math.min(slides.length-1,i));render();if(list)slides[index].scrollIntoView();else window.scrollTo(0,0)}
prev.addEventListener('click',()=>go(index-1));next.addEventListener('click',()=>go(index+1));jump.addEventListener('change',()=>go(+jump.value));all.addEventListener('click',()=>{list=!list;render();window.scrollTo(0,0)});document.querySelector('#print').addEventListener('click',()=>window.print());document.addEventListener('keydown',e=>{if(list||/INPUT|SELECT|TEXTAREA|BUTTON|A/.test(e.target.tagName)||e.ctrlKey||e.altKey||e.metaKey)return;const keys={ArrowRight:index+1,ArrowLeft:index-1,Home:0,End:slides.length-1};if(e.key in keys){e.preventDefault();go(keys[e.key])}});window.addEventListener('hashchange',()=>go(fromHash()));document.querySelector('.toolbar').hidden=false;render();})();'''
parts=['<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Tema de tesis MDS · Mario Romero</title><style>'+css+'</style></head><body>'+nav+'<main id="contenido" tabindex="-1">']
prs=Presentation();prs.slide_width=Pt(1280);prs.slide_height=Pt(720)
prs.core_properties.title='Tema de tesis MDS · Mario Romero'
prs.core_properties.author='Mario Romero';prs.core_properties.language='es-CL'
audit={'revision_date':'2026-09-23','slides':len(slides),'editable':True,'slides_verified':[]}
for n,s in enumerate(slides,1):
    parts.append(f'<section class="slide" aria-labelledby="s{n}" data-title="{html.escape(s["title"] if n>1 else "Tema y antecedentes",quote=True)}">')
    svg=[];texts=[];ps=prs.slides.add_slide(prs.slide_layouts[6])
    for obj in s['objects']:
        k=obj['kind']
        if k=='rect':
            svg.append(f'<rect x="{obj["x"]}" y="{obj["y"]}" width="{obj["w"]}" height="{obj["h"]}" fill="#{obj["fill"]}" stroke="#{obj["stroke"] or obj["fill"]}"/>')
            sh=ps.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(obj['x']),Pt(obj['y']),Pt(obj['w']),Pt(obj['h']))
            sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(obj['fill'])
            if obj['stroke']:sh.line.color.rgb=RGBColor.from_string(obj['stroke']);sh.line.width=Pt(1.3)
            else:sh.line.fill.background()
            sh._element.spPr.append(OxmlElement('a:effectLst'))
        elif k=='line':
            svg.append(f'<line x1="{obj["x1"]}" y1="{obj["y1"]}" x2="{obj["x2"]}" y2="{obj["y2"]}" stroke="#{obj["color"]}" stroke-width="{obj["width"]}"/>')
            sh=ps.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(obj['x1']),Pt(obj['y1']),Pt(obj['x2']),Pt(obj['y2']))
            sh.line.color.rgb=RGBColor.from_string(obj['color']);sh.line.width=Pt(obj['width'])
            sh._element.spPr.append(OxmlElement('a:effectLst'))
        else:
            tag=obj['tag'];attr=f' id="s{n}"' if obj['value']==s['title'] else ''
            style=f'left:{obj["x"]/12.8}%;top:{obj["y"]/7.2}%;width:{obj["w"]/12.8}%;height:{obj["h"]/7.2}%;font-size:{obj["size"]/12.8}cqw;font-family:{obj["font"]},serif;font-weight:{700 if obj["bold"] else 400};color:#{obj["color"]}'
            texts.append(f'<{tag}{attr} class="t {obj["cls"]}" style="{style}">{html.escape(obj["value"])}</{tag}>')
            if obj['mobile_after']:
                texts.append('<p class="mobile-only">'+html.escape(obj['mobile_after'])+'</p>')
            sh=ps.shapes.add_textbox(Pt(obj['x']),Pt(obj['y']),Pt(obj['w']),Pt(obj['h']))
            tf=sh.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            for i,value in enumerate(obj['value'].split('\n')):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=value
                p.font.name=obj['font'];p.font.size=Pt(obj['size']);p.font.bold=obj['bold'];p.font.color.rgb=RGBColor.from_string(obj['color'])
                p.space_before=p.space_after=Pt(0);p.line_spacing=1.1
    parts.append('<svg class="graphic" viewBox="0 0 1280 720" aria-hidden="true">'+''.join(svg)+'</svg>'+''.join(texts)+'</section>')
    ps.notes_slide.notes_text_frame.text='\n\n'.join(o['value'] for o in s['objects'] if o['kind']=='text')
    outside=[sh.shape_id for sh in ps.shapes if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width or sh.top+sh.height>prs.slide_height]
    assert not outside,(n,outside)
    audit['slides_verified'].append({'slide':n,'text_boxes':sum(o['kind']=='text' for o in s['objects']),'out_of_bounds':outside})
parts.append('</main><footer class="footer">Flechas para navegar · “Ver todas” para leer de corrido · Funciona sin conexión</footer><script>'+js+'</script></body></html>')
DEST.with_suffix('.html').write_text('\n'.join(parts),encoding='utf-8')
prs.save(DEST.with_suffix('.pptx'))
loaded=Presentation(DEST.with_suffix('.pptx'))
for s,ps in zip(slides,loaded.slides):
    values=[sh.text for sh in ps.shapes if sh.has_text_frame and sh.text]
    expected=[o['value'] for o in s['objects'] if o['kind']=='text']
    assert values==expected
with zipfile.ZipFile(DEST.with_suffix('.pptx')) as z:
    assert not [p for p in z.namelist() if p.startswith('ppt/media/')]
    assert not [p for p in z.namelist() if p.endswith('.rels') and b'TargetMode="External"' in z.read(p)]
audit.update(source_sha256=hashlib.sha256(DEST.with_suffix('.html').read_bytes()).hexdigest(),pptx_sha256=hashlib.sha256(DEST.with_suffix('.pptx').read_bytes()).hexdigest(),source_text_preserved=True,objectives_match_thesis_notes=True)
(ROOT/'documentacion/Verificacion_PPTX_Tema_23-09-2026.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'Created {len(slides)} slides: HTML and editable PPTX.')
