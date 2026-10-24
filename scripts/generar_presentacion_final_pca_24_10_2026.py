"""Extiende la presentación visual vigente con la auditoría PCA v2."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "presentaciones/Tema-Tesis-KDD-Visual-24-09-2026.pptx"
OUTPUT = ROOT / "presentaciones/Presentacion-Final-Tesis-PCA-24-10-2026.pptx"

NAVY = RGBColor(22, 50, 79)
TEAL = RGBColor(6, 106, 112)
BLUE = RGBColor(20, 93, 160)
GRAY = RGBColor(75, 89, 104)
LINE = RGBColor(205, 215, 223)
PALE = RGBColor(241, 245, 248)
WHITE = RGBColor(255, 255, 255)
AMBER = RGBColor(154, 92, 21)


def text_box(slide, x, y, w, h, content, size=18, color=NAVY, bold=False,
             font="Arial", align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = content
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, fill=PALE, border=LINE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(0.7)
    return shape


def new_slide(prs, kicker, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    box(slide, 0.47, 0.43, 0.28, 0.035, TEAL, TEAL)
    text_box(slide, 0.88, 0.30, 11.1, 0.3, kicker, 11, GRAY)
    text_box(slide, 0.48, 0.72, 12.2, 0.70, title, 29, NAVY, font="Cambria")
    box(slide, 0.48, 7.08, 12.35, 0.012, LINE, LINE)
    text_box(slide, 0.49, 7.12, 11.8, 0.22,
             "Fuente: Auditoria_PCA_Senales_v2_24-10-2026.md · análisis exploratorio", 9, GRAY)
    text_box(slide, 12.52, 7.10, 0.35, 0.24, str(number), 10, GRAY, align=PP_ALIGN.RIGHT)
    return slide


def add_audit(prs):
    slide = new_slide(prs, "Revisión del pipeline · 24-10-2026", "La auditoría cambió la referencia de evaluación", 8)
    items = [
        ("Alineación histórica", "21.719/21.719 ventanas válidas tenían el componente pupilar en una fila incorrecta.", NAVY),
        ("Profesor GSR v1", "El PC1 cargaba sobre dos variables SCR que valían cero en ~97 % de train.", NAVY),
        ("Contexto temporal", "El promedio cruzaba 28 huecos entre ventanas elegibles.", NAVY),
    ]
    for i, (heading, body, color) in enumerate(items):
        x = 0.48 + i * 4.17
        box(slide, x, 1.80, 3.95, 2.72)
        text_box(slide, x + 0.20, 2.03, 3.55, 0.55, heading, 18, TEAL, True)
        text_box(slide, x + 0.20, 2.71, 3.55, 1.49, body, 17, color)
    box(slide, 0.49, 5.04, 12.30, 1.23, WHITE, TEAL)
    text_box(slide, 0.77, 5.16, 11.8, 0.86,
             "v2: normalización con ventanas elegibles, GSR continuo y reinicio del historial ante huecos.",
             19, NAVY, True)
    text_box(slide, 0.53, 6.53, 12.1, 0.35,
             "Las métricas anteriores quedan como registro histórico; no son comparables con la v2.",
             14, AMBER)


def add_pipeline(prs):
    slide = new_slide(prs, "Transformación y minería de datos", "Un PCA por señal, en dos etapas separadas", 9)
    box(slide, 0.49, 1.69, 5.96, 4.77)
    box(slide, 6.77, 1.69, 6.04, 4.77)
    text_box(slide, 0.76, 1.87, 5.42, 0.45, "Modelo Profesor · crea etiquetas", 20, TEAL, True)
    text_box(slide, 0.78, 2.56, 5.34, 0.61, "Atención", 17, NAVY, True)
    text_box(slide, 0.78, 3.02, 5.25, 0.68, "Mirada 4 → 1 PC (64,2 %)\nPupila 1 → 1 PC", 17, NAVY)
    text_box(slide, 0.78, 4.04, 5.20, 0.49, "Activación", 17, NAVY, True)
    text_box(slide, 0.78, 4.50, 5.20, 0.74, "GSR tónica + fásica 2 → 1 PC (52,1 %)\nContexto causal de 6 s", 17, NAVY)
    text_box(slide, 0.78, 5.72, 5.26, 0.41, "Terciles por participante; protocolo offline", 13, GRAY)

    text_box(slide, 7.04, 1.87, 5.46, 0.45, "Modelo Estudiante · predice", 20, TEAL, True)
    text_box(slide, 7.04, 2.56, 5.28, 0.61, "Atención", 17, NAVY, True)
    text_box(slide, 7.04, 3.02, 5.20, 0.68, "EEG 9 → 4 PC + GSR 9 → 5 PC\n18 entradas → 9 componentes", 17, NAVY)
    text_box(slide, 7.04, 4.04, 5.20, 0.49, "Activación", 17, NAVY, True)
    text_box(slide, 7.04, 4.50, 5.20, 0.74, "EEG 9 → 4 + mirada 10 → 7 + pupila 6 → 4\n25 entradas → 15 componentes", 17, NAVY)
    text_box(slide, 7.04, 5.72, 5.29, 0.41, "Mínimo para retener ≥90 % por señal", 13, GRAY)
    text_box(slide, 0.54, 6.64, 12.10, 0.32,
             "Los PCA, imputadores y modelos se ajustan solo con train; validation y test se transforman.",
             14, BLUE)


def add_results(prs):
    slide = new_slide(prs, "Evaluación · 4.017 ventanas y ocho participantes en test",
                      "Reducir entradas conserva casi toda la predicción", 10)
    headings = [("Objetivo", 0.67, 2.18), ("Completo", 3.04, 1.61),
                ("PCA ≥90 %", 4.93, 1.70), ("1 PC/señal", 6.96, 1.81),
                ("Solo estímulo", 9.16, 2.30)]
    for title, x, w in headings:
        text_box(slide, x, 1.80, w, 0.45, title, 16, TEAL, True)
    rows = [
        ("Atención", "0,3893", "0,3871", "0,3834", "0,4610"),
        ("Activación", "0,3637", "0,3635", "0,3757", "0,4167"),
    ]
    for i, row in enumerate(rows):
        y = 2.40 + i * 1.13
        box(slide, 0.52, y, 12.25, 0.88, PALE if i == 0 else WHITE, LINE)
        for value, (_, x, w) in zip(row, headings):
            text_box(slide, x, y + 0.17, w, 0.49, value, 20 if x > 0.67 else 18,
                     NAVY if x < 9.16 else AMBER, x < 9.16)
    box(slide, 0.52, 4.94, 12.25, 1.25, WHITE, TEAL)
    text_box(slide, 0.78, 5.10, 11.77, 0.84,
             "El PCA ≥90 % reduce casi a la mitad las entradas. Solo estímulo predice mejor: la validez de las etiquetas sigue abierta.",
             18, NAVY, True)
    text_box(slide, 0.55, 6.57, 12.05, 0.34,
             "BA = precisión balanceada; Dummy = 0,3333. Test exploratorio, previamente observado.",
             13, GRAY)


def add_retry(prs):
    slide = new_slide(prs, "Reintentos de predicción · control por estímulo",
                      "La mejora en validation no se sostuvo en test", 11)
    headings = [("Tarea", 0.72, 2.1), ("Solo estímulo", 3.24, 2.2),
                ("Estímulo + PCA", 6.22, 2.65), ("Cambio", 10.12, 1.6)]
    for heading, x, w in headings:
        text_box(slide, x, 1.76, w, 0.48, heading, 16, TEAL, True)
    rows = [("Atención", "0,4613", "0,4608", "−0,0005"),
            ("Activación", "0,4120", "0,3766", "−0,0353")]
    for i, row in enumerate(rows):
        y = 2.38 + i * 1.05
        box(slide, 0.52, y, 12.25, 0.83)
        for value, (_, x, w) in zip(row, headings):
            text_box(slide, x, y + 0.17, w, 0.47, value, 20, NAVY if x < 10 else AMBER,
                     x < 10)
    box(slide, 0.54, 4.65, 12.20, 1.44, WHITE, TEAL)
    text_box(slide, 0.80, 4.78, 11.65, 1.14,
             "En validation, el PCA parecía sumar: atención 0,4728 frente a 0,4659; activación 0,4325 frente a 0,4082. El resultado no se generalizó.",
             18, NAVY, True)
    text_box(slide, 0.56, 6.35, 12.10, 0.56,
             "Un PC/señal dio 0,4658 en atención, pero 0,4722 con las señales desordenadas. Test exploratorio ya consultado.",
             13, AMBER)


def main() -> None:
    prs = Presentation(SOURCE)
    if len(prs.slides) != 7:
        raise ValueError("La presentación base cambió; revisar antes de agregar resultados")
    add_audit(prs)
    add_pipeline(prs)
    add_results(prs)
    add_retry(prs)
    prs.save(OUTPUT)
    print(f"{OUTPUT}: {len(prs.slides)} láminas")


if __name__ == "__main__":
    main()
